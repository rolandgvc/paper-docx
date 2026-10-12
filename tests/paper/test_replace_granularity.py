"""Word granularity for `Span.replace`: revisions that start and end on word boundaries."""

from __future__ import annotations

import datetime as dt
from pathlib import Path
from typing import List

import pytest

import docx
from docx.document import Document
from docx.oxml.ns import qn
from docx.search import find_one

from .harness.contract import assert_changed_parts, save_and_reopen

FROZEN = dt.datetime(2026, 1, 2, 3, 4, 5, tzinfo=dt.timezone.utc)


def _saved(tmp_path: Path, text: str) -> Path:
    document = docx.Document()
    document.add_paragraph(text)
    original = tmp_path / "original.docx"
    document.save(str(original))
    return original


def _revision_text(document: Document, tag: str) -> List[str]:
    body = document.element.body
    if tag == "w:del":
        return [
            "".join(t.text or "" for t in d.iter(qn("w:delText"))) for d in body.iter(qn("w:del"))
        ]
    return ["".join(t.text or "" for t in i.iter(qn("w:t"))) for i in body.iter(qn("w:ins"))]


class DescribeReplaceGranularity:
    def it_keeps_character_alignment_by_default(self, tmp_path: Path):
        original = _saved(tmp_path, "Term. Two (2) years from the Effective Date.")
        document = docx.Document(str(original))

        find_one(document, "Two (2) years").replace(
            "Twelve (12) months", tracked=True, author="Reviewer", date=FROZEN
        )
        reopened = save_and_reopen(document, tmp_path / "out.docx")

        assert _revision_text(reopened, "w:del") == ["o (2) year"]
        assert _revision_text(reopened, "w:ins") == ["elve (12) month"]

    def it_redlines_whole_words(self, tmp_path: Path):
        original = _saved(tmp_path, "Term. Two (2) years from the Effective Date.")
        document = docx.Document(str(original))

        find_one(document, "Two (2) years").replace(
            "Twelve (12) months",
            tracked=True,
            author="Reviewer",
            date=FROZEN,
            granularity="word",
        )
        out = tmp_path / "out.docx"
        reopened = save_and_reopen(document, out)

        assert _revision_text(reopened, "w:del") == ["Two (2) years"]
        assert _revision_text(reopened, "w:ins") == ["Twelve (12) months"]
        assert_changed_parts(original, out, ["word/document.xml"])

    def it_still_changes_only_the_word_that_differs(self, tmp_path: Path):
        original = _saved(tmp_path, "See Section 3 below.")
        document = docx.Document(str(original))

        find_one(document, "Section 3").replace(
            "Section 4", tracked=True, author="Reviewer", date=FROZEN, granularity="word"
        )
        reopened = save_and_reopen(document, tmp_path / "out.docx")

        assert _revision_text(reopened, "w:del") == ["3"]
        assert _revision_text(reopened, "w:ins") == ["4"]

    def it_widens_an_edit_inside_a_word_to_the_whole_word(self, tmp_path: Path):
        original = _saved(tmp_path, "The payment is due.")
        document = docx.Document(str(original))

        find_one(document, "payment").replace(
            "payments", tracked=True, author="Reviewer", date=FROZEN, granularity="word"
        )
        reopened = save_and_reopen(document, tmp_path / "out.docx")

        assert _revision_text(reopened, "w:del") == ["payment"]
        assert _revision_text(reopened, "w:ins") == ["payments"]

    def it_applies_to_untracked_replacements(self, tmp_path: Path):
        original = _saved(tmp_path, "Term. Two (2) years from the Effective Date.")
        document = docx.Document(str(original))

        result = find_one(document, "Two (2) years").replace(
            "Twelve (12) months", granularity="word"
        )
        reopened = save_and_reopen(document, tmp_path / "out.docx")

        assert reopened.paragraphs[0].text == "Term. Twelve (12) months from the Effective Date."
        assert result.preserved_formatting_regions

    @pytest.mark.parametrize("tracked", [False, True])
    @pytest.mark.parametrize("prefix", ["", "The "])
    @pytest.mark.parametrize("suffix", [" is due.", " is d"])
    def it_keeps_word_boundaries_across_formatting_regions(
        self, tmp_path: Path, tracked: bool, prefix: str, suffix: str
    ):
        document = docx.Document()
        paragraph = document.add_paragraph(prefix)
        paragraph.add_run("pay").bold = True
        paragraph.add_run("ment" + suffix).italic = True
        paragraph.add_run("ue." if suffix == " is d" else "")

        find_one(document, prefix + "payment is due.").replace(
            prefix + "payments is due.",
            tracked=tracked,
            author="Reviewer" if tracked else None,
            date=FROZEN,
            granularity="word",
        )
        reopened = save_and_reopen(document, tmp_path / "out.docx")

        if tracked:
            assert "".join(_revision_text(reopened, "w:del")) == "payment"
            assert _revision_text(reopened, "w:ins") == ["payments"]
        else:
            assert reopened.paragraphs[0].text == prefix + "payments is due."
        suffix_runs = [run for run in reopened.paragraphs[0].runs if run.text == suffix]
        assert len(suffix_runs) == 1
        assert suffix_runs[0].italic is True

    def it_refuses_an_unknown_granularity_before_changing_anything(self):
        document = docx.Document()
        document.add_paragraph("Two (2) years")
        before = document.element.xml

        with pytest.raises(ValueError, match="granularity"):
            find_one(document, "Two (2) years").replace(
                "Twelve (12) months",
                granularity="sentence",  # type: ignore[arg-type]
            )
        assert document.element.xml == before
